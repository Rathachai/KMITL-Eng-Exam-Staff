# -*- coding: utf-8 -*-
"""รายงานตามผู้คุมสอบ -> reports/by-staffs.xlsx"""

import openpyxl

from common import *   # pd, np, re, os, datetime, config values, helpers

_state = load_state()
dfX, dfST, dfslot = _state["dfX"], _state["dfST"], _state["dfslot"]

########################################################################
# EXPORT BY STAFF
########################################################################


# ─────────────────────────────── EXPORT BY STAFF ─────────────────────────────
import re

STAFF_COLS    = ['staff', 'ex_date', 'ex_dur', 'room', 'building']
STAFF_COLS_TH = ['ผู้คุมสอบ', 'วันที่', 'เวลา', 'ห้อง', 'อาคาร']

# ── Column alignment map ──────────────────────────────────────────────────
CENTER_COLS = {'วันที่', 'เวลา'}

# ── Title stripping (key for matching names with the master file) ─────────
TITLE_PAT = re.compile(
    r'^(ศ\.\s*ดร\.|รศ\.\s*ดร\.|ผศ\.\s*ดร\.|ศ\.|รศ\.|ผศ\.|ดร\.|อ\.|'
    r'นางสาว|นาง|นาย|น\.ส\.)\s*'
)

def strip_title(s) -> str:
    """'อ. อมตะ อนันต์พินิจวัฒนา' → 'อมตะ อนันต์พินิจวัฒนา'"""
    s = ' '.join(str(s).split())
    prev = None
    while prev != s:                       # stacked titles e.g. 'ผศ. ดร.'
        prev = s
        s = TITLE_PAT.sub('', s).strip()
    return s


def format_thai_date_short(d) -> str:
    """date → 'พฤ. 19 มี.ค. 69'"""
    dt = pd.Timestamp(d)
    return f"{THAI_DAYS_SHORT[dt.weekday()]} {dt.day} {THAI_MONTHS_SHORT[dt.month]} {(dt.year + 543) % 100}"


def build_header_lines_staff(dept: str) -> list[tuple[str, str]]:
    sem_num, sem_year = exam_semester(dfX['ex_date'].dropna().min()).split('/')
    etype             = exam_type_th(extype)
    return [
        ("สถาบันเทคโนโลยีพระจอมเกล้าเจ้าคุณทหารลาดกระบัง", "center"),
        (f"รายงานการจัดสอบ{etype}",                          "center"),
        (f"ประจำภาคเรียนที่ {sem_num} ปีการศึกษา {sem_year}", "center"),
        ("คณะวิศวกรรมศาสตร์",                               "center"),
        (f"หน่วยงาน: {dept}",                                "center"),
        ("",                                                  "center"),
    ]


def make_staff_sheet(ws, df_slot: pd.DataFrame, dept: str):
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    ncols       = len(df_slot.columns)
    header_rows = build_header_lines_staff(dept)
    n_header    = len(header_rows)

    thin    = Side(style='thin', color='000000')
    no_side = Side(style=None)

    # ── Header text block ─────────────────────────────────────────────────────
    for i, (text, align) in enumerate(header_rows, start=1):
        if ncols > 1:
            ws.merge_cells(start_row=i, start_column=1,
                           end_row=i,   end_column=ncols)
        cell           = ws.cell(row=i, column=1, value=text)
        cell.font      = Font(name='TH Sarabun New', bold=True, size=14)
        cell.alignment = Alignment(horizontal=align, vertical='center')

    # ── Column headers ────────────────────────────────────────────────────────
    col_row  = n_header + 1
    hdr_font = Font(name='TH Sarabun New', bold=True, size=14, color='000000')
    hdr_fill = PatternFill('solid', start_color='FFFFFF')
    for c, name in enumerate(df_slot.columns, start=1):
        cell           = ws.cell(row=col_row, column=c, value=name)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.border    = Border(left=thin, right=thin, top=thin, bottom=thin)
        cell.alignment = Alignment(horizontal='center', vertical='center',
                                   wrap_text=False)
    ws.row_dimensions[col_row].height = 24

    # ── Pre-compute last row of each staff group ──────────────────────────────
    staff_col       = 'ผู้คุมสอบ'
    staff_last_rows = set()
    if staff_col in df_slot.columns:
        orig_staff = df_slot[staff_col].tolist()
        for i, name in enumerate(orig_staff):
            is_last = (i == len(orig_staff) - 1) or (orig_staff[i + 1] != name)
            if is_last:
                staff_last_rows.add(col_row + 1 + i)

    # ── Deduplicate: show staff name only on first row of each group ──────────
    df_out = df_slot.copy()
    if staff_col in df_out.columns:
        seen = set()
        for idx in df_out.index:
            s = df_out.at[idx, staff_col]
            if s in seen:
                df_out.at[idx, staff_col] = ''
            else:
                seen.add(s)

    # ── Column index map ──────────────────────────────────────────────────────
    col_idx = {name: i + 1 for i, name in enumerate(df_out.columns)}
    staff_c = col_idx.get(staff_col)

    def make_border(excel_row: int, col: int) -> Border:
        if col == staff_c:
            is_last = excel_row in staff_last_rows
            return Border(left=thin, right=thin, top=no_side,
                          bottom=thin if is_last else no_side)
        return Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── Data rows ─────────────────────────────────────────────────────────────
    for r, row in enumerate(df_out.itertuples(index=False), start=col_row + 1):
        for c, value in enumerate(row, start=1):
            col_name = df_out.columns[c - 1]
            if isinstance(value, list):
                value = ', '.join(str(v) for v in value)
            cell           = ws.cell(row=r, column=c, value=value)
            cell.font      = Font(name='TH Sarabun New', size=13)
            cell.border    = make_border(r, c)
            cell.alignment = Alignment(
                horizontal = 'center' if col_name in CENTER_COLS else 'left',
                wrap_text  = False,
                vertical   = 'top'
            )

    # ── Column widths ─────────────────────────────────────────────────────────
    col_widths = {'ผู้คุมสอบ': 30, 'วันที่': 16, 'เวลา': 15, 'ห้อง': 12, 'อาคาร': 20}
    for c, name in enumerate(df_out.columns, start=1):
        ws.column_dimensions[get_column_letter(c)].width = col_widths.get(name, 14)

    # ── Print setup ───────────────────────────────────────────────────────────
    from openpyxl.worksheet.page import PageMargins
    ws.page_setup.orientation  = 'portrait'
    ws.page_setup.fitToPage    = True
    ws.page_setup.fitToWidth   = 1
    ws.page_setup.fitToHeight  = 0
    ws.page_setup.paperSize    = ws.PAPERSIZE_A4
    ws.page_margins            = PageMargins(
        left=0.5, right=0.5, top=0.75, bottom=0.75, header=0.3, footer=0.3
    )
    ws.print_title_rows        = f'1:{col_row}'
    ws.oddFooter.center.text   = 'หน้า &P/&N'
    ws.oddFooter.center.size   = 12
    ws.oddFooter.center.font   = 'TH Sarabun New'


# ── Master lookup (title / role), keyed by name without title ─────────────────
dfST_k       = dfST.copy()
dfST_k.index = dfST_k.index.map(strip_title)
dfST_k       = dfST_k[~dfST_k.index.duplicated(keep='first')]

role_order = dfslot['sort_order'].to_dict()      # role → sort order

def staff_info(name):
    """Return (display_name, sort_order) for one staff name from dfX."""
    key = strip_title(name)
    if key in dfST_k.index:
        title = dfST_k.at[key, 'formal_title']
        title = '' if pd.isna(title) else str(title).strip()
        role  = dfST_k.at[key, 'role']
        return f"{title} {key}".strip(), role_order.get(role, 99)
    return ' '.join(str(name).split()), 99        # not in master: keep as is

# Diagnostics
names = dfX['staff'].dropna().unique()
n_ok  = sum(strip_title(n) in dfST_k.index for n in names)
print(f"master match: {n_ok} / {len(names)}")

# ── Build workbook ────────────────────────────────────────────────────────────
wb_staff = openpyxl.Workbook()
wb_staff.remove(wb_staff.active)

all_depts = sorted(dfX['st_group_dept'].dropna().unique())
if 'none' in all_depts:
    all_depts.remove('none')

for dept in all_depts:
    df_dept = dfX[dfX['st_group_dept'] == dept].copy()
    if df_dept.empty:
        continue

    info = df_dept['staff'].map(staff_info)
    df_dept['_display'] = info.map(lambda t: t[0])
    df_dept['_sort']    = info.map(lambda t: t[1])

    # order: role (sort_order) → name → exam date/time (as before)
    df_dept = (df_dept
               .sort_values(['_sort', '_display', 'ex_dt'])
               .reset_index(drop=True))

    df_dept['staff']   = df_dept['_display']
    df_dept['ex_date'] = df_dept['ex_date'].apply(format_thai_date_short)

    df_dept = df_dept[STAFF_COLS]
    df_dept.columns = STAFF_COLS_TH

    ws = wb_staff.create_sheet(title=dept[:31])
    make_staff_sheet(ws, df_dept, dept)

# Guard: Excel needs at least one visible sheet
if not wb_staff.sheetnames:
    ws = wb_staff.create_sheet("ไม่มีข้อมูล")
    ws['A1'] = "ไม่พบข้อมูลผู้คุมสอบ (ตรวจสอบ st_group_dept)"

out_staff = out_path('by-staffs.xlsx')
wb_staff.save(out_staff)
print(f"Saved: {out_staff}")


