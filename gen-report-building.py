# -*- coding: utf-8 -*-
"""รายงานตามอาคาร/เวลา -> reports/by-building-time.xlsx"""

import openpyxl

from common import *   # pd, np, re, os, datetime, config values, helpers

_state = load_state()
extype = _state["extype"]
dfX = _state["dfX"]

############################################################
#  EXPORT BY BUILDING & TIME ───────────────────
############################################################



# ─────────────────────────────── EXPORT BY BUILDING-TIME ─────────────────────

BLDG_COL_WIDTHS = {
    'เวลา': 20, 'รหัสวิชา': 14, 'ชื่อวิชา': 36, 'กลุ่ม': 7,
    'ห้องสอบ': 14, 'ผู้รับข้อสอบ': 16, 'ผู้ส่งข้อสอบ': 16,
    'ซอง': 8, 'รับข้อสอบไปตรวจ': 18, 'หมายเหตุ': 14,
}

BIG_COLS    = {'รหัสวิชา', 'ห้องสอบ'}     # larger + bold
CENTER_COLS = {'เวลา', 'รหัสวิชา', 'ห้องสอบ'}   # horizontally centered



def build_header_lines_bldg(building: str, session: str,
                             ex_date, extype: str) -> list[tuple[str, str]]:
    dt      = pd.Timestamp(ex_date)
    day_s   = THAI_DAYS_SHORT[dt.weekday()]
    month_s = THAI_MONTHS_SHORT[dt.month]
    year_be = dt.year + 543
    sem     = exam_semester(ex_date)
    etype   = exam_type_th(extype)
    time_s  = SESSION_DURATION.get('09' if session == 'เช้า' else '13', '')

    return [
        ("สถาบันเทคโนโลยีพระจอมเกล้าเจ้าคุณทหารลาดกระบัง", "center"),
        (f"ตารางรับ-ส่งข้อสอบ  คณะวิศวกรรมศาสตร์ การสอบ{etype} ภาคการศึกษาที่ {sem}", "center"),
        (f"อาคาร{building}   วัน{day_s}ที่ {dt.day} {month_s} {year_be} ช่วง{session}   เวลา {time_s}", "center"),
    ]


def make_bldg_sheet(ws, df_slot: pd.DataFrame, building: str,
                    session: str, ex_date, extype: str):
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    ncols       = len(df_slot.columns)
    header_rows = build_header_lines_bldg(building, session, ex_date, extype)
    n_header    = len(header_rows)

    thin  = Side(style='thin',   color='000000')
    thick = Side(style='medium', color='000000')     # use 'thick' for even bolder

    # ── Header text block ─────────────────────────────────────────────────────
    for i, (text, align) in enumerate(header_rows, start=1):
        if ncols > 1:
            ws.merge_cells(start_row=i, start_column=1,
                           end_row=i,   end_column=ncols)
        cell           = ws.cell(row=i, column=1, value=text)
        cell.font      = Font(name='TH Sarabun New', bold=True, size=16)
        cell.alignment = Alignment(horizontal=align, vertical='center')
        ws.row_dimensions[i].height = 28

    # ── Blank row between header and column headers ───────────────────────────
    blank_row = n_header + 1
    ws.row_dimensions[blank_row].height = 10

    # ── Column headers (thick top + bottom border) ────────────────────────────
    col_row  = n_header + 2
    hdr_font = Font(name='TH Sarabun New', bold=True, size=14, color='000000')
    hdr_fill = PatternFill('solid', start_color='FFFFFF')
    hdr_bord = Border(left=thin, right=thin, top=thick, bottom=thick)
    for c, name in enumerate(df_slot.columns, start=1):
        cell           = ws.cell(row=col_row, column=c, value=name)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.border    = hdr_bord
        cell.alignment = Alignment(horizontal='center', vertical='center',
                                   wrap_text=False)
    ws.row_dimensions[col_row].height = 28

    # ── Last row of each room group (thick bottom border) ─────────────────────
    data_start_row = col_row + 1
    rooms          = df_slot['ห้องสอบ'].tolist()
    room_last_rows = {
        data_start_row + i
        for i, rm in enumerate(rooms)
        if i == len(rooms) - 1 or rooms[i + 1] != rm
    }

    # ── Data rows ─────────────────────────────────────────────────────────────
    ROW_HEIGHT = 52
    for r, row in enumerate(df_slot.itertuples(index=False), start=data_start_row):
        ws.row_dimensions[r].height = ROW_HEIGHT
        border = Border(
            left=thin, right=thin,
            top=thick if r == data_start_row else thin,   # match header's thick bottom
            bottom=thick if r in room_last_rows else thin,
        )
        for c, value in enumerate(row, start=1):
            col_name = df_slot.columns[c - 1]
            if isinstance(value, list):
                value = ', '.join(str(v) for v in value)
            cell        = ws.cell(row=r, column=c, value=value)
            cell.border = border
            cell.font   = (Font(name='TH Sarabun New', size=20, bold=True)
                           if col_name in BIG_COLS
                           else Font(name='TH Sarabun New', size=14))
            cell.alignment = Alignment(
                horizontal = 'center' if col_name in CENTER_COLS else None,
                vertical   = 'center',
                wrap_text  = (col_name == 'ชื่อวิชา'),
            )

    # ── Column widths ─────────────────────────────────────────────────────────
    for c, name in enumerate(df_slot.columns, start=1):
        ws.column_dimensions[get_column_letter(c)].width = BLDG_COL_WIDTHS.get(name, 12)

    # ── Print setup ───────────────────────────────────────────────────────────
    from openpyxl.worksheet.page import PageMargins
    ws.page_setup.orientation  = 'landscape'
    ws.page_setup.fitToPage    = True
    ws.page_setup.fitToWidth   = 1
    ws.page_setup.fitToHeight  = 0
    ws.page_setup.paperSize    = ws.PAPERSIZE_A4
    ws.page_margins            = PageMargins(
        left=0.5, right=0.5, top=0.75, bottom=0.75, header=0.3, footer=0.3
    )
    ws.print_title_rows        = f'1:{col_row}'
    ws.oddFooter.right.text    = 'หน้า &P/&N'
    ws.oddFooter.right.size    = 12
    ws.oddFooter.right.font    = 'TH Sarabun New'


# ── Prepare subject columns ───────────────────────────────────────────────────
dfX['subj_code'] = dfX['subj'].str.split().str[0].str.strip()
dfX['subj_name'] = dfX['subj'].str.split().str[1:].str.join(' ').str.strip()
dfX['is_extra']  = ~dfX['subj_code'].str.match(r'^\d+$')

# ── Build workbook ────────────────────────────────────────────────────────────
wb_bldg = openpyxl.Workbook()
wb_bldg.remove(wb_bldg.active)

all_dates     = sorted(dfX['ex_date'].dropna().unique())
all_buildings = sorted(dfX['building'].dropna().unique())

for ex_date in all_dates:
    df_date = dfX[dfX['ex_date'] == ex_date]

    for hour_key, session, session_en in [('09', 'เช้า', 'am'), ('13', 'บ่าย', 'pm')]:
        df_s = df_date[df_date['ex_time'] == hour_key]
        if df_s.empty:
            continue

        for building in all_buildings:
            df_day = df_s[df_s['building'] == building].copy()
            df_day = df_day[~df_day['is_extra']]
            if df_day.empty:
                continue

            df_day = df_day.sort_values(['room', 'ex_dur'])

            dt         = pd.Timestamp(ex_date)
            month_en   = THAI_MONTHS_SHORT[dt.month].replace('.', '')
            day_str    = f"{dt.day:02d}-{month_en}-{session}"
            sheet_name = f"{day_str} {building}"[:31]

            df_out = pd.DataFrame({
                'เวลา':             df_day['ex_dur'].values,
                'รหัสวิชา':         df_day['subj_code'].values,
                'ชื่อวิชา':         df_day['subj_name'].values,
                'กลุ่ม':            df_day['sec'].values,
                'ห้องสอบ':          df_day['room'].values,
                'ผู้รับข้อสอบ':     '',
                'ผู้ส่งข้อสอบ':     '',
                'ซอง':              '',
                'รับข้อสอบไปตรวจ': '',
                'หมายเหตุ':         '',
            }).reset_index(drop=True)

            ws = wb_bldg.create_sheet(title=sheet_name)
            make_bldg_sheet(ws, df_out, building, session, ex_date, extype)

# Guard: Excel needs at least one visible sheet
if not wb_bldg.sheetnames:
    ws = wb_bldg.create_sheet("ไม่มีข้อมูล")
    ws['A1'] = "ไม่พบข้อมูลการสอบ"

out_bldg = out_path('by-building-time.xlsx')
wb_bldg.save(out_bldg)
print(f"Saved: {out_bldg}")






