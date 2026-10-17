# -*- coding: utf-8 -*-
"""รายงานตามวิชา -> reports/by-subjects.xlsx"""

import openpyxl

from common import *   # pd, np, re, os, datetime, config values, helpers

_state = load_state()
dfX, dtlist = _state["dfX"], _state["dtlist"]

########################################################################
# EXPORT PER-SLOT EXCEL 
########################################################################



EXPORT_COLS    = ['order', 'subj', 'sec', 'st_year', 'st_max', 'ex_dur',
                  'teacher_str', 'building', 'room', 'note', 'staff']
EXPORT_COLS_TH = ['ลำดับ', 'วิชา', 'กลุ่ม', 'ชั้นปี', 'นศ.', 'เวลา',
                  'ผู้สอน', 'อาคาร', 'ห้อง', 'หมายเหตุ', 'ผู้คุมสอบ']


def build_header_lines(ex_dt_str: str) -> list[tuple[str, str]]:
    """
    Return list of (text, alignment) for the header block.
    alignment: 'center' | 'right'
    All rows are bold.
    """
    dt, hour_str = parse_ex_dt(ex_dt_str)
    day_s    = THAI_DAYS_SHORT[dt.weekday()]
    month_s  = THAI_MONTHS_SHORT[dt.month]
    year_be  = (dt.year + 543) % 100
    session  = SESSION_DURATION.get(hour_str, f'{hour_str}:00 น.')
    date_line = f"วันสอบ  {day_s} {dt.day} {month_s} {year_be}   เวลา {session}"

    return [
        ("สถาบันเทคโนโลยีพระจอมเกล้าเจ้าคุณทหารลาดกระบัง", "center"),
        ("รายงานการจัดสอบปลายภาค",                          "center"),
        ("ประจำภาคเรียนที่ 2 ปีการศึกษา 2568",              "center"),
        ("คณะวิศวกรรมศาสตร์",                               "center"),
        ("",                                                  "center"),
        (date_line,                                           "right"),
    ]


def make_slot_sheet(ws, df_slot: pd.DataFrame, ex_dt_str: str):
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    ncols       = len(df_slot.columns)
    header_rows = build_header_lines(ex_dt_str)
    n_header    = len(header_rows)

    thin    = Side(style='thin', color='000000')
    no_side = Side(style=None)

    # ── Header text block ─────────────────────────────────────────────────────
    for i, (text, align) in enumerate(header_rows, start=1):
        if ncols > 1:
            ws.merge_cells(start_row=i, start_column=1,
                           end_row=i,   end_column=ncols)
        cell           = ws.cell(row=i, column=1, value=text)
        cell.font      = Font(name='TH Sarabun New', bold=True, size=14)
        cell.alignment = Alignment(horizontal=align, vertical='center')

    # ── Column headers ────────────────────────────────────────────────────────
    col_row  = n_header + 1
    hdr_font = Font(name='TH Sarabun New', bold=True, size=14, color='000000')
    hdr_fill = PatternFill('solid', start_color='FFFFFF')
    for c, name in enumerate(df_slot.columns, start=1):
        cell           = ws.cell(row=col_row, column=c, value=name)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.border    = Border(left=thin, right=thin, top=thin, bottom=thin)
        cell.alignment = Alignment(horizontal='center', vertical='center',
                                   wrap_text=False)
    ws.row_dimensions[col_row].height = 24

    # ── Deduplicate: blank subj/order for repeated subjects ───────────────────
    subj_col  = 'วิชา'
    order_col = 'ลำดับ'
    df_out = df_slot.copy()
    if subj_col in df_out.columns:
        seen = set()
        for idx in df_out.index:
            subj = df_out.at[idx, subj_col]
            if subj in seen:
                df_out.at[idx, subj_col]  = ''
                if order_col in df_out.columns:
                    df_out.at[idx, order_col] = ''
            else:
                seen.add(subj)

    # ── Pre-compute last row of each subject group ────────────────────────────
    subj_last_rows = set()
    if subj_col in df_slot.columns:
        orig_subj = df_slot[subj_col].tolist()
        for i, subj in enumerate(orig_subj):
            is_last = (i == len(orig_subj) - 1) or (orig_subj[i + 1] != subj)
            if is_last:
                subj_last_rows.add(col_row + 1 + i)

    # ── Column index map ──────────────────────────────────────────────────────
    col_idx = {name: i + 1 for i, name in enumerate(df_out.columns)}
    subj_c  = col_idx.get(subj_col)
    order_c = col_idx.get(order_col)

    def make_border(excel_row: int, col: int) -> Border:
        if col in (subj_c, order_c):
            if excel_row in subj_last_rows:
                return Border(left=thin, right=thin, top=no_side, bottom=thin)
            else:
                return Border(left=thin, right=thin, top=no_side, bottom=no_side)
        return Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── Data rows ─────────────────────────────────────────────────────────────
    for r, row in enumerate(df_out.itertuples(index=False), start=col_row + 1):
        for c, value in enumerate(row, start=1):
            if isinstance(value, list):
                value = ', '.join(str(v) for v in value)
            is_subj        = (c == subj_c)
            cell           = ws.cell(row=r, column=c, value=value)
            cell.font      = Font(name='TH Sarabun New', size=14 if is_subj else 13)
            cell.border    = make_border(r, c)
            cell.alignment = Alignment(wrap_text=False, vertical='top')

    # ── Column widths ─────────────────────────────────────────────────────────
    col_widths = {
        'ลำดับ': 7, 'วิชา': 38, 'กลุ่ม': 8, 'ชั้นปี': 8,
        'นศ.': 7, 'เวลา': 16, 'ผู้สอน': 28,
        'อาคาร': 10, 'ห้อง': 10, 'หมายเหตุ': 14, 'ผู้คุมสอบ': 22,
    }
    for c, name in enumerate(df_out.columns, start=1):
        ws.column_dimensions[get_column_letter(c)].width = col_widths.get(name, 12)


    # ── Print setup: fit to page width ───────────────────────────────────────
    from openpyxl.worksheet.page import PageMargins, PrintPageSetup

    ws.page_setup.orientation     = 'landscape'
    ws.page_setup.fitToPage       = True
    ws.page_setup.fitToWidth      = 1
    ws.page_setup.fitToHeight     = 0   # 0 = ไม่จำกัดความสูง (ขยายลงได้กี่หน้าก็ได้)
    ws.page_setup.paperSize       = ws.PAPERSIZE_A4

    ws.page_margins = PageMargins(
        left=0.5, right=0.5, top=0.75, bottom=0.75,
        header=0.3, footer=0.3
    )

    # ── Repeat header rows on every page ─────────────────────────────────────
    # แถว 1 ถึง col_row (header block + หัว column)
    ws.print_title_rows = f'1:{col_row}'

    # ── Footer: page number ───────────────────────────────────────────────────
    ws.oddFooter.center.text = 'หน้า &P/&N'
    ws.oddFooter.center.size = 12
    ws.oddFooter.center.font = 'TH Sarabun New'


# ── Build workbook ────────────────────────────────────────────────────────────
import openpyxl

dfX2 = dfX[EXPORT_COLS].copy()
dfX2.columns = EXPORT_COLS_TH

wb_subj = openpyxl.Workbook()
wb_subj.remove(wb_subj.active)

for dt in sorted(dtlist):                          # เรียงตามวันเวลา
    df_slot = dfX2[dfX['ex_dt'] == dt].sort_values('วิชา').reset_index(drop=True)
    if df_slot.empty:
        continue

    ws = wb_subj.create_sheet(title=slot_sheet_name(dt))
    make_slot_sheet(ws, df_slot, dt)

out_path = out_path('by-subjects.xlsx')
wb_subj.save(out_path)
print(f"Saved: {out_path}")


