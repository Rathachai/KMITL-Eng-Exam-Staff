# -*- coding: utf-8 -*-
"""รายงานตามห้องสอบ -> reports/by-room.xlsx"""

import openpyxl

from common import *   # pd, np, re, os, datetime, config values, helpers

_state = load_state()
dfX, dtlist = _state["dfX"], _state["dtlist"]

########################################################################
# EXPORT BY ROOM
########################################################################

# ─────────────────────────────── EXPORT BY ROOM ──────────────────────────────

ROOM_COLS    = ['room', 'subj', 'st_num', 'staff']
ROOM_COLS_TH = ['ห้อง', 'วิชา', 'จำนวน', 'ผู้คุมสอบ']


def build_header_lines_room(ex_dt_str: str, building: str) -> list[tuple[str, str]]:
    dt, hour_str = parse_ex_dt(ex_dt_str)
    day_s    = THAI_DAYS_SHORT[dt.weekday()]
    month_s  = THAI_MONTHS_SHORT[dt.month]
    year_be  = (dt.year + 543) % 100
    session  = SESSION_DURATION.get(hour_str, f'{hour_str}:00 น.')
    date_line = f"อาคาร {building}   วันสอบ  {day_s} {dt.day} {month_s} {year_be}   เวลา {session}"

    return [
        ("สถาบันเทคโนโลยีพระจอมเกล้าเจ้าคุณทหารลาดกระบัง", "center"),
        ("รายงานการจัดสอบปลายภาค",                          "center"),
        ("ประจำภาคเรียนที่ 2 ปีการศึกษา 2568",              "center"),
        ("คณะวิศวกรรมศาสตร์",                               "center"),
        ("",                                                  "center"),
        (date_line,                                           "right"),
    ]


def make_room_sheet(ws, df_slot: pd.DataFrame, ex_dt_str: str, building: str):
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    ncols       = len(df_slot.columns)
    header_rows = build_header_lines_room(ex_dt_str, building)
    n_header    = len(header_rows)

    thin    = Side(style='thin', color='000000')
    no_side = Side(style=None)

    # ── Header text block ─────────────────────────────────────────────────────
    for i, (text, align) in enumerate(header_rows, start=1):
        if ncols > 1:
            ws.merge_cells(start_row=i, start_column=1,
                           end_row=i,   end_column=ncols)
        cell           = ws.cell(row=i, column=1, value=text)
        cell.font      = Font(name='TH Sarabun New', bold=True, size=14)
        cell.alignment = Alignment(horizontal=align, vertical='center')

    # ── Column headers ────────────────────────────────────────────────────────
    col_row  = n_header + 1
    hdr_font = Font(name='TH Sarabun New', bold=True, size=14, color='000000')
    hdr_fill = PatternFill('solid', start_color='FFFFFF')
    for c, name in enumerate(df_slot.columns, start=1):
        cell           = ws.cell(row=col_row, column=c, value=name)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.border    = Border(left=thin, right=thin, top=thin, bottom=thin)
        cell.alignment = Alignment(horizontal='center', vertical='center',
                                   wrap_text=False)
    ws.row_dimensions[col_row].height = 24

    # ── Pre-compute last row of each room group ───────────────────────────────
    room_col      = 'ห้อง'
    room_last_rows = set()
    if room_col in df_slot.columns:
        orig_room = df_slot[room_col].tolist()
        for i, room in enumerate(orig_room):
            is_last = (i == len(orig_room) - 1) or (orig_room[i + 1] != room)
            if is_last:
                room_last_rows.add(col_row + 1 + i)

    # ── Deduplicate room: show only on first row of each group ────────────────
    df_out = df_slot.copy()
    if room_col in df_out.columns:
        seen = set()
        for idx in df_out.index:
            room = df_out.at[idx, room_col]
            if room in seen:
                df_out.at[idx, room_col] = ''
            else:
                seen.add(room)

    # ── Column index map ──────────────────────────────────────────────────────
    col_idx  = {name: i + 1 for i, name in enumerate(df_out.columns)}
    room_c   = col_idx.get('ห้อง')
    subj_c   = col_idx.get('วิชา')
    staff_c  = col_idx.get('ผู้คุมสอบ')

    def make_border(excel_row: int, col: int) -> Border:
        is_last = excel_row in room_last_rows
        if col in (room_c, subj_c):
            # เส้นขอบล่างเฉพาะแถวสุดท้ายของห้อง
            return Border(left=thin, right=thin, top=no_side,
                          bottom=thin if is_last else no_side)
        # ผู้คุมสอบ และ column อื่น: full box ทุกแถว
        return Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── Data rows ─────────────────────────────────────────────────────────────
    for r, row in enumerate(df_out.itertuples(index=False), start=col_row + 1):
        for c, value in enumerate(row, start=1):
            if isinstance(value, list):
                value = ', '.join(str(v) for v in value)
            is_subj        = (df_out.columns[c - 1] == 'วิชา')
            cell           = ws.cell(row=r, column=c, value=value)
            cell.font      = Font(name='TH Sarabun New', size=14 if is_subj else 13)
            cell.border    = make_border(r, c)
            cell.alignment = Alignment(wrap_text=False, vertical='top')

    # ── Column widths ─────────────────────────────────────────────────────────
    col_widths = {'ห้อง': 12, 'วิชา': 60, 'จำนวน':10, 'ผู้คุมสอบ': 40}
    for c, name in enumerate(df_out.columns, start=1):
        ws.column_dimensions[get_column_letter(c)].width = col_widths.get(name, 14)

    # ── Print setup ───────────────────────────────────────────────────────────
    from openpyxl.worksheet.page import PageMargins
    ws.page_setup.orientation  = 'landscape'
    ws.page_setup.fitToPage    = True
    ws.page_setup.fitToWidth   = 1
    ws.page_setup.fitToHeight  = 0
    ws.page_setup.paperSize    = ws.PAPERSIZE_A4
    ws.page_margins            = PageMargins(
        left=0.5, right=0.5, top=0.75, bottom=0.75, header=0.3, footer=0.3
    )
    ws.print_title_rows        = f'1:{col_row}'
    ws.oddFooter.center.text   = 'หน้า &P/&N'
    ws.oddFooter.center.size   = 12
    ws.oddFooter.center.font   = 'TH Sarabun New'


# ── Build workbook ────────────────────────────────────────────────────────────
wb_room = openpyxl.Workbook()
wb_room.remove(wb_room.active)

all_buildings = sorted(dfX['building'].dropna().unique())

for building in all_buildings:
    for dt in sorted(dtlist):
        df_dt = dfX[(dfX['ex_dt'] == dt) & (dfX['building'] == building)].copy()
        if df_dt.empty:
            continue

        df_b = (df_dt.sort_values(['room', 'subj'])[ROOM_COLS]
                .reset_index(drop=True))
        df_b.columns = ROOM_COLS_TH

        sheet_name = f"{building} {slot_sheet_name(dt)}"[:31]
        ws = wb_room.create_sheet(title=sheet_name)
        make_room_sheet(ws, df_b, dt, building)

out_room = out_path('by-room.xlsx')
wb_room.save(out_room)
print(f"Saved: {out_room}")        



