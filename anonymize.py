# -*- coding: utf-8 -*-
"""
anonymize.py — แปลงชื่อบุคคลในไฟล์ใน exam-data/ เป็นรหัสตามภาค/ตำแหน่ง
เขียนผลลงโฟลเดอร์ใหม่ exam-data-anon/ (ไม่แก้ต้นฉบับ)

    ชื่อใหม่ = <ภาค>-<ตำแหน่ง>-<เลขลำดับ>   เช่น  คอม-อจ-001 (ภาคภาษาอังกฤษใช้ตัวพิมพ์เล็ก), คอม-ผบรห-002
    คนที่ไม่อยู่ใน master (อาจารย์ภายนอก ฯลฯ)    =  นอก-อจ-001

คำนำหน้า (รศ. ดร. นางสาว ฯลฯ) คงไว้ ส่วนชื่อ-นามสกุลถูกแทนที่ทุกเซลล์ในทุกชีต
ตารางแปลง (ชื่อจริง -> รหัส) เขียนที่ anon-mapping.csv  ห้าม commit (อยู่ใน .gitignore)

    python3 anonymize.py
"""

import csv
import glob
import os
import re

import openpyxl
import pandas as pd

from common import clean_name_list, DATA_DIR, MASTER_FILE, BASE_DIR

OUT = os.path.join(BASE_DIR, 'exam-data-anon')
MAPPING_CSV = os.path.join(BASE_DIR, 'anon-mapping.csv')

# main_dept -> รหัสย่อ (แก้ได้) ; ภาคที่ไม่อยู่ในนี้ใช้ชื่อเดิมของ main_dept
DEPT_CODE = {
    'วิศวกรรมไฟฟ้า': 'ไฟฟ้า',
    'วิศวกรรมโยธา': 'โยธา',
    'วิศวกรรมเครื่องกล': 'กล',
    'วิศวกรรมอุตสาหการ': 'อุต',
    'วิศวกรรมการวัดและควบคุม': 'วัด',
    'วิศวกรรมคอมพิวเตอร์': 'คอม',
    'วิศวกรรมโทรคมนาคม': 'โทร',
    'วิศวกรรมอิเล็กทรอนิกส์': 'อิเล็ก',
    'วิศวกรรมเคมี': 'เคมี',
    'วิศวกรรมชีวการแพทย์': 'ชีว',
    'วิศวกรรมเกษตร': 'เกษตร',
    'วิศวกรรมไอโอทีและสารสนเทศ': 'ไอโอที',
    'วิศวกรรมอาหาร': 'อาหาร',
    'วิศวกรรมหุ่นยนต์และปัญญาประดิษฐ์': 'หุ่น',
    'สำนักงานบริหารหลักสูตรวิศวกรรมสหวิทยาการนานาชาติ (SIIE)': 'siie',
    'ศูนย์นวัตกรรมและการถ่ายทอดเทคโนโลยีปัญญาประดิษฐ์ CiRE CORE (SECAC)': 'cire',
}
ROLE_CODE = {
    'อาจารย์': 'อจ',
    'อาจารย์ (ประเภทพิเศษ)': 'อจพ',
    'ผู้บริหาร': 'ผบรห',
    'หัวหน้าภาค': 'หนภ',
    'หัวหน้างาน': 'หนง',
    'เจ้าหน้าที่': 'จนท',
}
EXTERNAL = 'นอก'

NAME_COLS_SCHEDULE = ('ผู้สอน', 'ผู้คุมสอบ')


def norm(s):
    return re.sub(r'\s+', ' ', str(s)).strip().lower()


def main():
    os.makedirs(OUT, exist_ok=True)
    master_src = os.path.join(DATA_DIR, MASTER_FILE)
    sched_srcs = sorted(glob.glob(os.path.join(DATA_DIR, 'ร*.xlsx')))

    # ── 1) master: สร้างรหัสให้ทุกคน ─────────────────────────────────────────
    st = pd.read_excel(master_src, sheet_name='exam_staff')
    counters, mapping, per_row = {}, {}, []
    for _, r in st.iterrows():
        full = norm(f"{'' if pd.isna(r['first']) else r['first']} {'' if pd.isna(r['last']) else r['last']}")
        dept = DEPT_CODE.get(r['main_dept'], str(r['main_dept']).strip()).lower()   # ตัวพิมพ์เล็ก: โค้ดจับคู่ชื่อแบบ lower
        role = ROLE_CODE.get(r['role'], str(r['role']).strip())
        if full in mapping:                       # ชื่อซ้ำใน master -> รหัสเดิม
            per_row.append(mapping[full])
            continue
        counters[(dept, role)] = counters.get((dept, role), 0) + 1
        mapping[full] = f"{dept}-{role}-{counters[(dept, role)]:03d}"
        per_row.append(mapping[full])
    n_master = len(mapping)

    # ── 2) ชื่อใน ตารางสอบ ที่ไม่อยู่ใน master -> นอก-อจ-### ─────────────────
    externals = {}
    for f in sched_srcs:
        df = pd.read_excel(f, header=1)
        for col in NAME_COLS_SCHEDULE:
            if col not in df.columns:
                continue
            for cell in df[col].dropna().astype(str):
                for nm in clean_name_list(cell.split(',')):
                    nm = norm(nm)
                    if nm and nm not in ('nan', '-') and nm not in mapping:
                        externals[nm] = externals.get(nm, 0) + 1
    ext_n = 0
    for nm in sorted(externals):
        ext_n += 1
        mapping[nm] = f"{EXTERNAL}-{ROLE_CODE['อาจารย์']}-{ext_n:03d}"

    # ── 3) regex แทนที่ชื่อเต็มทุกเซลล์ (ยาวก่อน) ───────────────────────────
    names = sorted(mapping, key=len, reverse=True)
    pat = re.compile('|'.join(r'\s+'.join(re.escape(t) for t in n.split()) for n in names),
                     re.IGNORECASE)

    def sub(text):
        return pat.sub(lambda m: mapping[norm(m.group(0))], text)

    def scrub(path_in, path_out, master=False):
        wb = openpyxl.load_workbook(path_in)
        wb.properties.creator = ''
        wb.properties.lastModifiedBy = ''
        for ws in wb.worksheets:
            first_col = last_col = None
            if master and ws.title == 'exam_staff':
                hdr = {c.value: c.column for c in ws[1]}
                first_col, last_col = hdr['first'], hdr['last']
            for row in ws.iter_rows():
                for c in row:
                    if not isinstance(c.value, str) or type(c).__name__ == 'MergedCell':
                        continue
                    if first_col and c.row > 1 and c.column in (first_col, last_col):
                        continue                    # เขียนเองด้านล่าง
                    c.value = sub(c.value)
            if first_col:
                for i, code in enumerate(per_row):
                    ws.cell(row=i + 2, column=first_col).value = code
                    ws.cell(row=i + 2, column=last_col).value = None
        wb.save(path_out)

    scrub(master_src, os.path.join(OUT, MASTER_FILE), master=True)
    for f in sched_srcs:
        scrub(f, os.path.join(OUT, os.path.basename(f)))

    with open(MAPPING_CSV, 'w', encoding='utf-8-sig', newline='') as fp:
        w = csv.writer(fp)
        w.writerow(['original', 'code'])
        w.writerows(sorted(mapping.items(), key=lambda kv: kv[1]))

    print(f"master: {n_master} คน | นอก master: {ext_n} คน | ไฟล์ตารางสอบ: {len(sched_srcs)}")
    print(f"เขียนที่ {OUT}")
    print(f"ตารางแปลง (ห้าม commit): {MAPPING_CSV}")


if __name__ == '__main__':
    main()
