# -*- coding: utf-8 -*-
"""ใบรับ-ส่งข้อสอบ (sign sheet) แยกตามวัน-เวลา-อาคาร -> reports/for-pickup.xlsx"""

import openpyxl

from common import *   # pd, np, re, os, datetime, config values, helpers

_state = load_state()
extype = _state["extype"]
dfX = _state["dfX"]

# ─────────────────────────────── EXPORT SIGN SHEET (รับ-ส่งข้อสอบ) ──────────

print('########## EXPORT SIGN SHEET (รับ-ส่งข้อสอบ)  ##########')


SIGN_COLS_TH = ['ห้องสอบ', 'รหัสวิชา', 'ชื่อวิชา', 'กลุ่มเรียน',
                'ผู้รับข้อสอบ', 'ผู้ส่งข้อสอบ', 'ผู้รับข้อสอบไปตรวจ', 'หมายเหตุ']


BIG_COLS      = {'ห้องสอบ', 'รหัสวิชา'}                    # larger + bold
CENTER_COLS   = {'ห้องสอบ', 'รหัสวิชา', 'กลุ่มเรียน'}      # horizontally centered
NAME_COL_SIZE = 16                                         # ชื่อวิชา: slightly larger + bold


def build_header_lines_sign(ex_date, session: str, building: str) -> list[tuple]:
    dt      = pd.Timestamp(ex_date)
    day_s   = THAI_DAYS_SHORT[dt.weekday()]
    month_s = THAI_MONTHS_SHORT[dt.month]
    year_be = dt.year + 543
    sem     = exam_semester(ex_date)
    etype   = exam_type_th(extype)
    time_s  = SESSION_DURATION.get('09' if session == 'เช้า' else '13', '')
    sem_num, sem_year = sem.split('/')
    date_line = f"วันสอบ  {day_s} {dt.day} {month_s} {year_be}   เวลา {time_s}"

    return [
        ("สถาบันเทคโนโลยีพระจอมเกล้าเจ้าคุณทหารลาดกระบัง",  "center"),
        (f"รายงานการจัดสอบ{etype}",                            "center"),
        (f"ประจำภาคเรียนที่ {sem_num} ปีการศึกษา {sem_year}",  "center"),
        ("คณะวิศวกรรมศาสตร์",                                  "center"),
        ((f"อาคาร {building}", date_line),                      "split"),  # building left, date right
    ]


def make_sign_sheet(ws, df_slot: pd.DataFrame, ex_date, session: str, building: str):
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    ncols       = len(df_slot.columns)
    header_rows = build_header_lines_sign(ex_date, session, building)
    n_header    = len(header_rows)

    thin    = Side(style='thin',   color='000000')
    thick   = Side(style='medium', color='000000')     # use 'thick' for even bolder
    no_side = Side(style=None)

    # ── Header text block ─────────────────────────────────────────────────────
    for i, (text, align) in enumerate(header_rows, start=1):
        ws.row_dimensions[i].height = 26

        if align == 'split':
            left_text, right_text = text
            mid = ncols // 2

            ws.merge_cells(start_row=i, start_column=1,
                           end_row=i,   end_column=mid)
            cell_l           = ws.cell(row=i, column=1, value=left_text)
            cell_l.font      = Font(name='TH Sarabun New', bold=True, size=16)
            cell_l.alignment = Alignment(horizontal='left', vertical='center')

            ws.merge_cells(start_row=i, start_column=mid + 1,
                           end_row=i,   end_column=ncols)
            cell_r           = ws.cell(row=i, column=mid + 1, value=right_text)
            cell_r.font      = Font(name='TH Sarabun New', bold=True, size=16)
            cell_r.alignment = Alignment(horizontal='right', vertical='center')

        else:
            if ncols > 1:
                ws.merge_cells(start_row=i, start_column=1,
                               end_row=i,   end_column=ncols)
            cell           = ws.cell(row=i, column=1, value=text)
            cell.font      = Font(name='TH Sarabun New', bold=True, size=16)
            cell.alignment = Alignment(horizontal=align, vertical='center')

    # ── Blank separator row ───────────────────────────────────────────────────
    blank_row = n_header + 1
    ws.row_dimensions[blank_row].height = 8

    # ── Column headers (thick top + bottom border) ────────────────────────────
    col_row  = n_header + 2
    hdr_font = Font(name='TH Sarabun New', bold=True, size=14, color='000000')
    hdr_fill = PatternFill('solid', start_color='FFFFFF')
    hdr_bord = Border(left=thin, right=thin, top=thick, bottom=thick)
    for c, name in enumerate(df_slot.columns, start=1):
        cell           = ws.cell(row=col_row, column=c, value=name)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.border    = hdr_bord
        cell.alignment = Alignment(horizontal='center', vertical='center',
                                   wrap_text=True)
    ws.row_dimensions[col_row].height = 28

    # ── Pre-compute last row of each room group ───────────────────────────────
    room_col       = 'ห้องสอบ'
    data_start_row = col_row + 1
    orig_room      = df_slot[room_col].tolist()
    room_last_rows = {
        data_start_row + i
        for i, room in enumerate(orig_room)
        if i == len(orig_room) - 1 or orig_room[i + 1] != room
    }

    # ── Deduplicate room column (merge-like) ──────────────────────────────────
    df_out = df_slot.copy()
    seen   = set()
    for idx in df_out.index:
        r = df_out.at[idx, room_col]
        if r in seen:
            df_out.at[idx, room_col] = ''
        else:
            seen.add(r)

    room_c = list(df_out.columns).index(room_col) + 1

    def make_border(excel_row: int, col: int) -> Border:
        is_last = excel_row in room_last_rows
        if col == room_c:
            # merge-like: no line between rows of the same room
            top = thick if excel_row == data_start_row else no_side
            return Border(left=thin, right=thin, top=top,
                          bottom=thick if is_last else no_side)
        # first data row touches the header's thick bottom line
        top = thick if excel_row == data_start_row else thin
        return Border(left=thin, right=thin, top=top,
                      bottom=thick if is_last else thin)

    # ── Data rows ─────────────────────────────────────────────────────────────
    ROW_HEIGHT = 48
    for r, row in enumerate(df_out.itertuples(index=False), start=data_start_row):
        ws.row_dimensions[r].height = ROW_HEIGHT
        for c, value in enumerate(row, start=1):
            col_name = df_out.columns[c - 1]
            if isinstance(value, list):
                value = ', '.join(str(v) for v in value)
            cell        = ws.cell(row=r, column=c, value=value)
            cell.border = make_border(r, c)

            if col_name in BIG_COLS:
                cell.font = Font(name='TH Sarabun New', size=20, bold=True)
            elif col_name == 'ชื่อวิชา':
                cell.font = Font(name='TH Sarabun New', size=NAME_COL_SIZE, bold=True)
            else:
                cell.font = Font(name='TH Sarabun New', size=14)

            cell.alignment = Alignment(
                horizontal = 'center' if col_name in CENTER_COLS else 'left',
                vertical   = 'center',
                wrap_text  = (col_name == 'ชื่อวิชา'),
            )

    # ── Column widths ─────────────────────────────────────────────────────────
    col_widths = {
        'ห้องสอบ': 14, 'รหัสวิชา': 14, 'ชื่อวิชา': 38, 'กลุ่มเรียน': 12,
        'ผู้รับข้อสอบ': 20, 'ผู้ส่งข้อสอบ': 20,
        'ผู้รับข้อสอบไปตรวจ': 22, 'หมายเหตุ': 16,
    }
    for c, name in enumerate(df_out.columns, start=1):
        ws.column_dimensions[get_column_letter(c)].width = col_widths.get(name, 14)

    # ── Print setup ───────────────────────────────────────────────────────────
    from openpyxl.worksheet.page import PageMargins
    ws.page_setup.orientation  = 'landscape'
    ws.page_setup.fitToPage    = True
    ws.page_setup.fitToWidth   = 1
    ws.page_setup.fitToHeight  = 0
    ws.page_setup.paperSize    = ws.PAPERSIZE_A4
    ws.page_margins            = PageMargins(
        left=0.5, right=0.5, top=0.75, bottom=0.75, header=0.3, footer=0.3
    )
    ws.print_title_rows        = f'1:{col_row}'
    ws.oddFooter.right.text    = 'หน้า &P/&N'
    ws.oddFooter.right.size    = 12
    ws.oddFooter.right.font    = 'TH Sarabun New'


# ── Build workbook ────────────────────────────────────────────────────────────
wb_sign = openpyxl.Workbook()
wb_sign.remove(wb_sign.active)

all_dates     = sorted(dfX['ex_date'].dropna().unique())
all_buildings = sorted(dfX['building'].dropna().unique())

for ex_date in all_dates:
    df_date = dfX[dfX['ex_date'] == ex_date]

    for hour_key, session, session_en in [('09', 'เช้า', 'am'), ('13', 'บ่าย', 'pm')]:
        df_s = df_date[df_date['ex_time'] == hour_key]
        if df_s.empty:
            continue

        for building in all_buildings:
            df_b = df_s[df_s['building'] == building].copy()
            if df_b.empty:
                continue

            df_b = df_b.sort_values(['room', 'subj']).reset_index(drop=True)

            dt         = pd.Timestamp(ex_date)
            month_en   = THAI_MONTHS_SHORT[dt.month].replace('.', '')
            day_str    = f"{dt.day:02d}-{month_en}-{session_en}"
            sheet_name = f"{day_str} {building}"[:31]   # day - time - building

            subj_parts = df_b['subj'].astype(str).str.split()

            df_out = pd.DataFrame({
                'ห้องสอบ':             df_b['room'].values,
                'รหัสวิชา':            subj_parts.str[0].values,
                'ชื่อวิชา':            subj_parts.str[1:].str.join(' ').values,
                'กลุ่มเรียน':          df_b['sec'].values,
                'ผู้รับข้อสอบ':        '',
                'ผู้ส่งข้อสอบ':        '',
                'ผู้รับข้อสอบไปตรวจ': '',
                'หมายเหตุ':            '',
            }).reset_index(drop=True)

            ws = wb_sign.create_sheet(title=sheet_name)
            make_sign_sheet(ws, df_out, ex_date, session, building)

# Guard: Excel needs at least one visible sheet
if not wb_sign.sheetnames:
    ws = wb_sign.create_sheet("ไม่มีข้อมูล")
    ws['A1'] = "ไม่พบข้อมูลการสอบ"

out_sign = out_path('for-pickup.xlsx')
wb_sign.save(out_sign)
print(f"Saved: {out_sign}")
